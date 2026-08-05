from __future__ import annotations

import html
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .models import InsightItem
from .sample_data import is_sample_url


KB_YELLOW = colors.HexColor("#FFCC00")
KB_DARK = colors.HexColor("#242424")
KB_GRAY = colors.HexColor("#5B5B5B")
KB_LIGHT = colors.HexColor("#F4F4F2")
KB_LINE = colors.HexColor("#D7D7D2")
KB_GOLD = colors.HexColor("#7A5B00")


def build_report_files(
    insights: List[InsightItem],
    output_dir: Path,
    metrics: Dict[str, Any],
    model_used: str,
    top_n: int = 5,
) -> Dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    now = datetime.now()
    date_compact = now.strftime("%Y%m%d")
    date_display = now.strftime("%Y.%m.%d")
    ranked = sorted(insights, key=lambda item: item.importance_score, reverse=True)[
        :top_n
    ]
    pdf_path = output_dir / f"Daily_Insight_{date_compact}.pdf"
    preview_path = output_dir / f"Daily_Insight_{date_compact}_preview.png"
    xlsx_path = output_dir / f"Daily_Insight_Items_{date_compact}.xlsx"
    subject = f"[Daily Insight] AI/IT 기술 동향 - {date_display}"
    html_body = build_email_html(ranked, metrics, date_display)

    build_pdf(
        ranked,
        pdf_path,
        metrics=metrics,
        model_used=model_used,
        generated_at=now,
    )
    build_preview_png(
        ranked,
        preview_path,
        metrics=metrics,
        generated_at=now,
    )
    build_xlsx(ranked, xlsx_path, generated_at=now)
    return {
        "pdf_path": pdf_path,
        "preview_path": preview_path,
        "xlsx_path": xlsx_path,
        "email_subject": subject,
        "html_body": html_body,
        "ranked": ranked,
    }


def build_email_html(
    insights: List[InsightItem],
    metrics: Dict[str, Any],
    date_display: str,
) -> str:
    categories = ", ".join(_top_categories(insights)) or "기술 동향"
    item_blocks = []
    for rank, item in enumerate(insights, start=1):
        keywords = " · ".join(html.escape(keyword) for keyword in item.keywords)
        source_link = (
            '<span style="font-size:12px;color:#777;">'
            "샘플 데이터 · 실제 원문 없음</span>"
            if is_sample_url(item.url)
            else (
                f'<a href="{html.escape(item.url, quote=True)}" '
                'style="font-size:12px;color:#624b00;">원문 보기</a>'
            )
        )
        item_blocks.append(
            f"""
            <tr>
              <td style="padding:20px 0;border-bottom:1px solid #deded8;">
                <div style="font-size:12px;color:#725800;margin-bottom:7px;">
                  TOP {rank} &nbsp;|&nbsp; {html.escape(item.category)}
                  &nbsp;|&nbsp; {item.importance_score}점
                </div>
                <div style="font-size:18px;font-weight:700;color:#222;margin-bottom:10px;">
                  {html.escape(item.title)}
                </div>
                <div style="font-size:14px;line-height:1.7;color:#333;margin-bottom:10px;">
                  {html.escape(item.summary)}
                </div>
                <div style="font-size:13px;line-height:1.6;color:#5b5b5b;margin-bottom:8px;">
                  <strong>실무 적용</strong> {html.escape(item.work_relevance)}
                </div>
                <div style="font-size:12px;color:#777;margin-bottom:8px;">
                  {html.escape(item.source)} · {html.escape(item.published_at)}
                  &nbsp;|&nbsp; {keywords}
                </div>
                {source_link}
              </td>
            </tr>
            """
        )
    return f"""<!doctype html>
<html lang="ko">
<body style="margin:0;background:#f2f2ef;font-family:Arial,'Apple SD Gothic Neo',sans-serif;color:#222;">
  <table role="presentation" width="100%" cellspacing="0" cellpadding="0">
    <tr>
      <td align="center" style="padding:24px 12px;">
        <table role="presentation" width="680" cellspacing="0" cellpadding="0"
               style="max-width:680px;background:#fff;border:1px solid #deded8;">
          <tr>
            <td style="height:10px;background:#ffcc00;"></td>
          </tr>
          <tr>
            <td style="padding:28px 32px 18px;">
              <div style="font-size:12px;color:#725800;">TECHLAB DAILY REPORT</div>
              <h1 style="font-size:28px;line-height:1.25;margin:8px 0 6px;">Daily Insight</h1>
              <div style="font-size:14px;color:#666;">AI/IT 기술 동향 · {date_display}</div>
            </td>
          </tr>
          <tr>
            <td style="padding:0 32px;">
              <table role="presentation" width="100%" cellspacing="0" cellpadding="0"
                     style="background:#f4f4f2;border-left:4px solid #ffcc00;">
                <tr>
                  <td style="padding:16px 18px;font-size:14px;line-height:1.6;color:#333;">
                    오늘은 <strong>{html.escape(categories)}</strong> 중심으로
                    사내 IT 개발 업무 관련성이 높은 {len(insights)}건을 선별했습니다.
                    원천 {metrics.get('raw_count', 0)}건 중 정제
                    {metrics.get('clean_count', 0)}건을 분석했습니다.
                  </td>
                </tr>
              </table>
            </td>
          </tr>
          <tr>
            <td style="padding:4px 32px 30px;">
              <table role="presentation" width="100%" cellspacing="0" cellpadding="0">
                {''.join(item_blocks)}
              </table>
            </td>
          </tr>
          <tr>
            <td style="padding:18px 32px;background:#242424;color:#cfcfc8;font-size:11px;line-height:1.6;">
              공개 외부 정보를 AI로 요약한 연구용 결과입니다.
              중요 의사결정 전 원문과 사내 보안 정책을 확인하십시오.
            </td>
          </tr>
        </table>
      </td>
    </tr>
  </table>
</body>
</html>"""


def build_pdf(
    insights: List[InsightItem],
    output_path: Path,
    metrics: Dict[str, Any],
    model_used: str,
    generated_at: datetime,
) -> None:
    font_name = _register_korean_font()
    document = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=20 * mm,
        bottomMargin=18 * mm,
        title="Daily Insight - AI/IT 기술 동향",
        author="TechLab",
        subject="AI Agent 기반 기술 동향 리포트",
    )
    styles = _pdf_styles(font_name)
    story: List[Any] = []
    date_display = generated_at.strftime("%Y.%m.%d")

    story.append(
        Table(
            [[""]],
            colWidths=[174 * mm],
            rowHeights=[5 * mm],
            style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), KB_YELLOW)]),
        )
    )
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("TECHLAB DAILY REPORT", styles["eyebrow"]))
    story.append(Paragraph("Daily Insight", styles["title"]))
    story.append(Paragraph(f"AI/IT 기술 동향 · {date_display}", styles["subtitle"]))
    story.append(Spacer(1, 8 * mm))

    top_categories = ", ".join(_top_categories(insights)) or "기술 동향"
    summary_text = (
        f"오늘은 {top_categories} 중심으로 사내 IT 개발 업무 관련성이 높은 "
        f"{len(insights)}건을 선별했습니다. 원천 {metrics.get('raw_count', 0)}건에서 "
        f"중복 {metrics.get('duplicate_count', 0)}건과 제외 "
        f"{metrics.get('excluded_count', 0)}건을 정리한 뒤 AI 분석을 수행했습니다."
    )
    summary_box = Table(
        [[Paragraph(_safe(summary_text), styles["body"])]],
        colWidths=[166 * mm],
        style=TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), KB_LIGHT),
                ("BOX", (0, 0), (-1, -1), 0.5, KB_LINE),
                ("LINEBEFORE", (0, 0), (0, -1), 4, KB_YELLOW),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ]
        ),
    )
    story.append(summary_box)
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("Top Insight", styles["section"]))

    for rank, item in enumerate(insights, start=1):
        keyword_text = " · ".join(item.keywords)
        score_line = (
            f"TOP {rank}  |  {item.category}  |  "
            f"{item.importance} {item.importance_score}점"
        )
        source_line = f"{item.source} · {item.published_at} · {keyword_text}"
        source_reference = (
            Paragraph(
                "샘플 데이터 · 실제 원문 링크 없음",
                styles["link"],
            )
            if is_sample_url(item.url)
            else Paragraph(
                f'<link href="{html.escape(item.url, quote=True)}" color="#6B5200">'
                f"{_safe(item.url)}</link>",
                styles["link"],
            )
        )
        content = [
            Paragraph(_safe(score_line), styles["rank"]),
            Paragraph(_safe(item.title), styles["item_title"]),
            Spacer(1, 2 * mm),
            Paragraph(_safe(item.summary), styles["body"]),
            Spacer(1, 2 * mm),
            Paragraph(
                _safe(f"실무 적용  {item.work_relevance}"),
                styles["work"],
            ),
            Spacer(1, 2 * mm),
            Paragraph(_safe(source_line), styles["meta"]),
            source_reference,
            Spacer(1, 2 * mm),
            Paragraph(_safe(f"주의  {item.risk_note}"), styles["risk"]),
            Spacer(1, 6 * mm),
        ]
        story.append(KeepTogether(content))

    story.append(PageBreak())
    story.append(Paragraph("처리 결과 및 검증 메모", styles["section"]))
    metric_rows = [
        ["항목", "결과"],
        ["원천 데이터", f"{metrics.get('raw_count', 0)}건"],
        ["정제 데이터", f"{metrics.get('clean_count', 0)}건"],
        ["중복 제거", f"{metrics.get('duplicate_count', 0)}건"],
        ["품질 기준 제외", f"{metrics.get('excluded_count', 0)}건"],
        ["AI Insight", f"{metrics.get('insight_count', len(insights))}건"],
        ["수집 성공률", f"{metrics.get('source_success_rate', 0)}%"],
        ["사용 모델", model_used],
    ]
    metric_table = Table(
        [
            [Paragraph(_safe(str(cell)), styles["table"]) for cell in row]
            for row in metric_rows
        ],
        colWidths=[52 * mm, 114 * mm],
        repeatRows=1,
    )
    metric_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), KB_DARK),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, KB_LINE),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.append(metric_table)
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("중요도 산정 기준", styles["section_small"]))
    story.append(
        Paragraph(
            "최신성 25점, 업무 관련성 30점, 확산 가능성 20점, "
            "금융권 연관성 15점, 출처 신뢰도 10점의 합계로 산정합니다.",
            styles["body"],
        )
    )
    story.append(Spacer(1, 5 * mm))
    story.append(Paragraph("사용 시 유의사항", styles["section_small"]))
    story.append(
        Paragraph(
            "이 리포트는 공개 외부 정보를 AI로 요약한 연구용 결과입니다. "
            "원문 왜곡 여부와 링크 유효성을 검수하고, 중요 의사결정 전에는 "
            "원문 및 사내 보안 정책을 확인해야 합니다.",
            styles["body"],
        )
    )

    document.build(
        story,
        onFirstPage=lambda canvas, doc: _draw_footer(canvas, doc, font_name),
        onLaterPages=lambda canvas, doc: _draw_footer(canvas, doc, font_name),
    )


def build_xlsx(
    insights: List[InsightItem],
    output_path: Path,
    generated_at: datetime,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Daily Insight"
    headers = [
        "date",
        "rank",
        "category",
        "importance",
        "score",
        "title",
        "source",
        "published_at",
        "url",
        "summary",
        "work_relevance",
        "risk_note",
        "keywords",
    ]
    sheet.append(headers)
    for rank, item in enumerate(insights, start=1):
        sheet.append(
            [
                generated_at.date().isoformat(),
                rank,
                item.category,
                item.importance,
                item.importance_score,
                item.title,
                item.source,
                item.published_at,
                (
                    "샘플 데이터 · 실제 원문 없음"
                    if is_sample_url(item.url)
                    else item.url
                ),
                item.summary,
                item.work_relevance,
                item.risk_note,
                ", ".join(item.keywords),
            ]
        )

    header_fill = PatternFill("solid", fgColor="242424")
    accent_fill = PatternFill("solid", fgColor="FFCC00")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet["A1"].fill = accent_fill
    sheet["A1"].font = Font(color="242424", bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    widths = [13, 8, 16, 10, 9, 38, 22, 15, 44, 70, 55, 42, 30]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(output_path)


def build_preview_png(
    insights: List[InsightItem],
    output_path: Path,
    metrics: Dict[str, Any],
    generated_at: datetime,
) -> None:
    width, height = 1240, 1754
    margin = 92
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    fonts = {
        "eyebrow": _image_font(20),
        "title": _image_font(54),
        "subtitle": _image_font(23),
        "section": _image_font(31),
        "rank": _image_font(18),
        "item_title": _image_font(30),
        "body": _image_font(22),
        "meta": _image_font(17),
    }

    draw.rectangle((margin, 72, width - margin, 93), fill="#FFCC00")
    draw.text(
        (margin, 132),
        "TECHLAB DAILY REPORT",
        font=fonts["eyebrow"],
        fill="#705500",
    )
    draw.text((margin, 172), "Daily Insight", font=fonts["title"], fill="#242424")
    draw.text(
        (margin, 248),
        f"AI/IT 기술 동향 · {generated_at.strftime('%Y.%m.%d')}",
        font=fonts["subtitle"],
        fill="#666660",
    )

    summary_top = 315
    summary_bottom = 430
    draw.rectangle(
        (margin, summary_top, width - margin, summary_bottom),
        fill="#F4F4F2",
        outline="#D7D7D2",
        width=1,
    )
    draw.rectangle(
        (margin, summary_top, margin + 10, summary_bottom),
        fill="#FFCC00",
    )
    categories = ", ".join(_top_categories(insights)) or "기술 동향"
    summary = (
        f"오늘은 {categories} 중심으로 사내 IT 개발 업무 관련성이 높은 "
        f"{len(insights)}건을 선별했습니다. 원천 {metrics.get('raw_count', 0)}건 중 "
        f"정제 {metrics.get('clean_count', 0)}건을 분석했습니다."
    )
    _draw_wrapped_text(
        draw,
        summary,
        (margin + 32, summary_top + 24),
        fonts["body"],
        "#333333",
        width - (margin * 2) - 62,
        line_gap=8,
        max_lines=3,
    )

    y = 482
    draw.text((margin, y), "Top Insight", font=fonts["section"], fill="#242424")
    y += 58
    for rank, item in enumerate(insights[:3], start=1):
        draw.text(
            (margin, y),
            f"TOP {rank} · {item.category} · {item.importance_score}점",
            font=fonts["rank"],
            fill="#705500",
        )
        y += 38
        y = _draw_wrapped_text(
            draw,
            item.title,
            (margin, y),
            fonts["item_title"],
            "#242424",
            width - (margin * 2),
            line_gap=7,
            max_lines=2,
        )
        y += 16
        y = _draw_wrapped_text(
            draw,
            item.summary,
            (margin, y),
            fonts["body"],
            "#343434",
            width - (margin * 2),
            line_gap=8,
            max_lines=3,
        )
        y += 18
        work_lines = _wrap_text(
            draw,
            f"실무 적용  {item.work_relevance}",
            fonts["body"],
            width - (margin * 2) - 42,
        )[:2]
        work_height = 25 + len(work_lines) * 34
        draw.rectangle(
            (margin, y, width - margin, y + work_height),
            fill="#F4F4F2",
        )
        draw.rectangle((margin, y, margin + 8, y + work_height), fill="#FFCC00")
        line_y = y + 15
        for line in work_lines:
            draw.text(
                (margin + 24, line_y),
                line,
                font=fonts["body"],
                fill="#343434",
            )
            line_y += 34
        y += work_height + 15
        meta = f"{item.source} · {item.published_at} · {' / '.join(item.keywords)}"
        y = _draw_wrapped_text(
            draw,
            meta,
            (margin, y),
            fonts["meta"],
            "#666660",
            width - (margin * 2),
            line_gap=5,
            max_lines=2,
        )
        y += 28
        draw.line((margin, y, width - margin, y), fill="#D7D7D2", width=1)
        y += 30
        if y > height - 160:
            break

    draw.line(
        (margin, height - 82, width - margin, height - 82),
        fill="#D7D7D2",
        width=1,
    )
    draw.text(
        (margin, height - 62),
        "TechLab Daily Insight · Research Demo",
        font=fonts["meta"],
        fill="#666660",
    )
    image.save(output_path, format="PNG", optimize=True)


def save_run_json(
    output_path: Path,
    payload: Dict[str, Any],
) -> None:
    with output_path.open("w", encoding="utf-8") as output_file:
        json.dump(payload, output_file, ensure_ascii=False, indent=2)


def _register_korean_font() -> str:
    font_candidates = [
        Path("/System/Library/Fonts/Supplemental/AppleGothic.ttf"),
        Path("/System/Library/Fonts/Supplemental/NotoSansGothic-Regular.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
        Path("C:/Windows/Fonts/malgun.ttf"),
    ]
    for font_path in font_candidates:
        if not font_path.exists():
            continue
        try:
            if "TechLabKorean" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("TechLabKorean", str(font_path)))
            return "TechLabKorean"
        except Exception:
            continue
    return "Helvetica"


def _image_font(size: int) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("/System/Library/Fonts/Supplemental/AppleGothic.ttf"),
        Path("/System/Library/Fonts/Supplemental/NotoSansGothic-Regular.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
        Path("C:/Windows/Fonts/malgun.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            try:
                return ImageFont.truetype(str(candidate), size=size)
            except OSError:
                continue
    return ImageFont.load_default()


def _draw_wrapped_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    position: tuple,
    font: ImageFont.ImageFont,
    fill: str,
    max_width: int,
    line_gap: int = 6,
    max_lines: Optional[int] = None,
) -> int:
    x, y = position
    lines = _wrap_text(draw, text, font, max_width)
    if max_lines:
        lines = lines[:max_lines]
    line_height = int(font.size * 1.35) if hasattr(font, "size") else 30
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill)
        y += line_height + line_gap
    return y


def _wrap_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    font: ImageFont.ImageFont,
    max_width: int,
) -> List[str]:
    words = " ".join(str(text).split()).split(" ")
    lines: List[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textlength(candidate, font=font) <= max_width:
            current = candidate
            continue
        if current:
            lines.append(current)
        current = word
    if current:
        lines.append(current)
    return lines or [""]


def _pdf_styles(font_name: str) -> Dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "eyebrow": ParagraphStyle(
            "Eyebrow",
            parent=base["Normal"],
            fontName=font_name,
            fontSize=8.5,
            leading=11,
            textColor=KB_GOLD,
            spaceAfter=3,
        ),
        "title": ParagraphStyle(
            "Title",
            parent=base["Title"],
            fontName=font_name,
            fontSize=27,
            leading=32,
            alignment=TA_LEFT,
            textColor=KB_DARK,
            spaceAfter=3,
        ),
        "subtitle": ParagraphStyle(
            "Subtitle",
            parent=base["Normal"],
            fontName=font_name,
            fontSize=10.5,
            leading=14,
            textColor=KB_GRAY,
        ),
        "section": ParagraphStyle(
            "Section",
            parent=base["Heading2"],
            fontName=font_name,
            fontSize=16,
            leading=21,
            textColor=KB_DARK,
            spaceBefore=2,
            spaceAfter=6,
        ),
        "section_small": ParagraphStyle(
            "SectionSmall",
            parent=base["Heading3"],
            fontName=font_name,
            fontSize=12,
            leading=16,
            textColor=KB_DARK,
            spaceAfter=5,
        ),
        "rank": ParagraphStyle(
            "Rank",
            parent=base["Normal"],
            fontName=font_name,
            fontSize=8.5,
            leading=11,
            textColor=KB_GOLD,
            spaceAfter=3,
        ),
        "item_title": ParagraphStyle(
            "ItemTitle",
            parent=base["Heading3"],
            fontName=font_name,
            fontSize=14,
            leading=19,
            textColor=KB_DARK,
            spaceAfter=2,
        ),
        "body": ParagraphStyle(
            "Body",
            parent=base["BodyText"],
            fontName=font_name,
            fontSize=9.4,
            leading=15,
            textColor=KB_DARK,
            wordWrap="CJK",
        ),
        "work": ParagraphStyle(
            "Work",
            parent=base["BodyText"],
            fontName=font_name,
            fontSize=9,
            leading=14,
            leftIndent=7,
            borderColor=KB_YELLOW,
            borderWidth=0,
            borderPadding=6,
            backColor=KB_LIGHT,
            textColor=KB_DARK,
            wordWrap="CJK",
        ),
        "meta": ParagraphStyle(
            "Meta",
            parent=base["Normal"],
            fontName=font_name,
            fontSize=8,
            leading=11,
            textColor=KB_GRAY,
            wordWrap="CJK",
        ),
        "link": ParagraphStyle(
            "Link",
            parent=base["Normal"],
            fontName=font_name,
            fontSize=7.5,
            leading=10,
            textColor=KB_GOLD,
            wordWrap="CJK",
        ),
        "risk": ParagraphStyle(
            "Risk",
            parent=base["Normal"],
            fontName=font_name,
            fontSize=8,
            leading=11,
            textColor=KB_GRAY,
            wordWrap="CJK",
        ),
        "table": ParagraphStyle(
            "Table",
            parent=base["Normal"],
            fontName=font_name,
            fontSize=8.5,
            leading=12,
            wordWrap="CJK",
        ),
    }


def _draw_footer(canvas: Any, document: Any, font_name: str) -> None:
    canvas.saveState()
    canvas.setFillColor(colors.white)
    canvas.rect(0, 0, A4[0], A4[1], fill=1, stroke=0)
    canvas.setStrokeColor(KB_LINE)
    canvas.setLineWidth(0.5)
    canvas.line(18 * mm, 13 * mm, 192 * mm, 13 * mm)
    canvas.setFont(font_name, 7)
    canvas.setFillColor(KB_GRAY)
    canvas.drawString(18 * mm, 8.5 * mm, "TechLab Daily Insight · Research Demo")
    canvas.drawRightString(
        192 * mm,
        8.5 * mm,
        f"{document.page}",
    )
    canvas.restoreState()


def _top_categories(insights: Iterable[InsightItem], limit: int = 3) -> List[str]:
    counter = Counter(item.category for item in insights)
    return [category for category, _ in counter.most_common(limit)]


def _safe(value: str) -> str:
    return html.escape(str(value)).replace("\n", "<br/>")
